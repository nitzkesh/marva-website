"""
Marva doc-template build script.

Produces (all in out\\, each with a Word-exported PDF):
  Marva-Report-HE.docx / .pdf   - Hebrew/RTL report (the APPROVED proof; byte/logic-stable)
  Marva-Report-EN.docx / .pdf   - English/LTR mirror
  Marva-Letter-HE.docx / .pdf   - Hebrew business letter
  Marva-Letter-EN.docx / .pdf   - English mirror
  Marva-Memo-HE.docx / .pdf     - Hebrew memo
  Marva-Memo-EN.docx / .pdf     - English mirror
  Marva-Quote-HE.docx / .pdf    - Hebrew price quote (NOT a tax invoice)
  Marva-Quote-EN.docx / .pdf    - English mirror
  assets/marva-mark.png         - rasterized full-color mark, transparent bg, ~600px wide
  assets/marva-mark-white.png   - rasterized reversed (off-white) mark, transparent bg

Reproducible / rerunnable: delete doc-templates/assets and doc-templates/out (or just
rerun) and this script regenerates everything from source SVGs + this file.

Source of truth for palette/type/logo: Marva/brand-book/Marva_brand_foundation.md

---------------------------------------------------------------------------------------
Architecture (read this before editing)
---------------------------------------------------------------------------------------
This is ONE parametrized generator over (doc_type x language):
  doc_type in {"report", "letter", "memo", "quote"}
  lang     in {"he", "en"}

Low-level oxml helpers (style_set, add_rtl_run, add_ltr_run, fix_doc_defaults, fix_theme,
_clean_rfonts, _set_cs_size/_set_cs_bold, ensure_fonts_loaded, add_page_field, ...) are
UNCHANGED from the original Hebrew-only proof script and are reused verbatim - they encode
hard-won Word/OOXML lessons (see each docstring) and are not doc-type or language specific.

Language-aware content is built through a small dispatch layer:
  body_font(lang)          - Assistant (he) / Work Sans (en)
  add_text(p, text, lang)  - routes to add_rtl_run (he) or add_ltr_run (en)
  start_paragraph(...)     - new paragraph, bidi-flagged only when he
  style_heading_runs(...)  - forces cs-font (+ rtl when he) on an already-created run
  set_style(...)           - style_set wrapper that only attaches the style-level
                              <w:rtl/> marker for he (see its docstring - this is the one
                              deliberate behavior fix needed to make the reused style_set()
                              safe for English styles)

Mirroring convention: HE is RTL with bidi paragraphs, JC omitted (start-of-reading edge
alignment, i.e. visually right); EN is plain LTR with NO bidi anywhere and default
(left) alignment. Header wordmark sits at reading-start, tagline at reading-end, in both
directions - achieved by toggling ONLY <w:bidiVisual/> on the header's 2-cell table and
the paragraph-alignment trick, never by reordering cell content. Footer contact line
(phone/email/site) is physically identical in both languages (contacts are Latin/digits),
built with add_ltr_run() directly regardless of doc language.

QA is direction- and doc-type-aware (see qa_checks/measure_pdf at the bottom): position
checks assert the right THIRD of the page for he and the left THIRD for en; table column
order checks assert bidiVisual mirroring for he and plain LTR order for en.
---------------------------------------------------------------------------------------
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
MARVA_ROOT = os.path.dirname(HERE)
ASSETS_DIR = os.path.join(HERE, "assets")
OUT_DIR = os.path.join(HERE, "out")

SVG_MARK = os.path.join(MARVA_ROOT, "website", "app", "public", "marva-mark.svg")
SVG_MARK_WHITE = os.path.join(MARVA_ROOT, "website", "app", "public", "marva-mark-white.svg")

PNG_MARK = os.path.join(ASSETS_DIR, "marva-mark.png")
PNG_MARK_WHITE = os.path.join(ASSETS_DIR, "marva-mark-white.png")

# kept for backward compatibility (default target = the approved HE report)
DOCX_OUT = os.path.join(OUT_DIR, "Marva-Report-HE.docx")
PDF_OUT = os.path.join(OUT_DIR, "Marva-Report-HE.pdf")

# ---------------------------------------------------------------------------
# Colors (from Marva_brand_foundation.md - use exactly, no others)
# ---------------------------------------------------------------------------
INK = "111111"
MUTED = "5B6660"
SAGE = "A8BFAE"
SAGE_TINT = "EEF2EC"

TARGET_PNG_WIDTH = 600  # px, target rasterization width

# ---------------------------------------------------------------------------
# RTL alignment convention (defect-4 fix).
#
# With <w:bidi/> set on a paragraph, Word interprets w:jc LOGICALLY:
# "left" means start-of-reading (visually RIGHT for a RTL paragraph) and
# "right" means end-of-reading (visually LEFT). So the original jc="right" +
# bidi combination rendered everything flush-left.
# JC_MODE options:
#   "omit"  - write no w:jc at all; the bidi paragraph default (start = right
#             edge) applies.
#   "left"  - write jc="left", which for a bidi paragraph = visual right.
#   "right" - write jc="right" (the original, broken convention).
# Empirical result (Word 16 PDF export, measured with pymupdf, 2026-07-11):
#   "omit"  -> title x1=514, header wordmark x1=501  (visual RIGHT, correct)
#   "left"  -> identical to "omit" (x1=514/501)      (visual RIGHT, correct)
#   "right" -> title x1=192, header wordmark x1=85   (visual LEFT, broken)
# LOCKED: "omit" - cleanest XML (no w:jc at all) and renders identically to
# "left". Use this convention for all future RTL templates.
# ---------------------------------------------------------------------------
JC_MODE = "omit"


def rasterize_svgs():
    """Rasterize the two mark SVGs to transparent PNGs at ~600px wide.

    Primary path: svglib + reportlab renderPM (backed by rlPyCairo).
    Fallback: cairosvg, if svglib fails on these files.
    """
    os.makedirs(ASSETS_DIR, exist_ok=True)

    def _via_svglib(svg_path, png_path):
        # svglib/reportlab's renderPM does not expose a true alpha channel for
        # PNG output (the 'transparent' configPIL key only applies to TIFF/PICT).
        # Work around it with difference matting: render the same drawing once
        # on pure white and once on pure black. For any pixel, if C_w and C_b
        # are the observed colors on white/black backgrounds:
        #   alpha   = 255 - mean(C_w - C_b)
        #   color   = C_b / (alpha/255)   (un-premultiply)
        # This recovers per-pixel alpha (incl. anti-aliased edges) without any
        # assumption about which colors appear in the artwork.
        from svglib.svglib import svg2rlg
        from reportlab.graphics import renderPM
        from PIL import Image
        import numpy as np

        def _render(bg_hex):
            drawing = svg2rlg(svg_path)
            if drawing is None:
                raise RuntimeError(f"svglib returned None for {svg_path}")
            scale = TARGET_PNG_WIDTH / drawing.width
            drawing.width *= scale
            drawing.height *= scale
            drawing.scale(scale, scale)
            tmp = png_path + f".tmp{bg_hex}.png"
            renderPM.drawToFile(drawing, tmp, fmt="PNG", bg=bg_hex, backendFmt="RGB")
            arr = np.array(Image.open(tmp).convert("RGB"), dtype=np.float64)
            os.remove(tmp)
            return arr

        on_white = _render(0xFFFFFF)
        on_black = _render(0x000000)

        diff = on_white - on_black  # >=0, ~0 where opaque, ~255 where background
        alpha = 255.0 - diff.mean(axis=2)
        alpha = np.clip(alpha, 0, 255)

        alpha_safe = np.where(alpha < 1, 1, alpha)  # avoid /0 on fully transparent px
        color = on_black / (alpha_safe[..., None] / 255.0)
        color = np.clip(color, 0, 255)

        rgba = np.dstack([color, alpha]).astype(np.uint8)
        Image.fromarray(rgba, mode="RGBA").save(png_path)

    def _via_cairosvg(svg_path, png_path):
        import cairosvg

        cairosvg.svg2png(url=svg_path, write_to=png_path, output_width=TARGET_PNG_WIDTH)

    for svg_path, png_path in ((SVG_MARK, PNG_MARK), (SVG_MARK_WHITE, PNG_MARK_WHITE)):
        try:
            _via_svglib(svg_path, png_path)
            print(f"[rasterize] svglib OK: {svg_path} -> {png_path}")
        except Exception as e_svglib:
            print(f"[rasterize] svglib failed for {svg_path}: {e_svglib}")
            try:
                _via_cairosvg(svg_path, png_path)
                print(f"[rasterize] cairosvg fallback OK: {svg_path} -> {png_path}")
            except Exception as e_cairo:
                raise RuntimeError(
                    f"Both svglib and cairosvg failed to rasterize {svg_path}.\n"
                    f"svglib error: {e_svglib}\ncairosvg error: {e_cairo}"
                )

        # Verify transparency (alpha channel present and not fully opaque)
        from PIL import Image

        img = Image.open(png_path)
        if img.mode != "RGBA":
            print(f"[rasterize][WARN] {png_path} is mode {img.mode}, not RGBA")
        else:
            alphas = img.getchannel("A").getextrema()
            if alphas[0] == 255:
                print(f"[rasterize][WARN] {png_path} alpha channel has no transparency (min alpha=255)")


# ---------------------------------------------------------------------------
# docx helpers (RTL / bidi / complex-script fonts / borders / fields need
# direct oxml manipulation - python-docx has no high-level API for them)
#
# Everything in this section is UNCHANGED from the original Hebrew-only proof
# script (byte-for-byte reused) except two additive, backward-compatible
# parameter defaults on fix_doc_defaults/fix_theme (see their docstrings).
# ---------------------------------------------------------------------------

RUBIK = "Rubik"
ASSISTANT = "Assistant"
WORK_SANS = "Work Sans"


def _rgb(hex_str):
    from docx.shared import RGBColor

    return RGBColor.from_string(hex_str)


def _clean_rfonts(rFonts, name):
    """Point an existing <w:rFonts> at `name` for ALL four slots and strip
    Word's theme attributes. In OOXML precedence the theme attributes
    (asciiTheme/hAnsiTheme/eastAsiaTheme/cstheme) WIN over the literal names,
    so leaving them in place silently resolves the font back to the theme
    (Calibri; majorBidi -> Times New Roman for Hebrew) even when the literal
    names say Rubik/Assistant. Round-2 root cause #1."""
    from docx.oxml.ns import qn

    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(attr)) is not None:
            del rFonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def _set_cs_size(rPr, size_pt):
    """Set w:szCs (complex-script size, half-points). Word sizes RTL/Hebrew
    text from w:szCs, not w:sz - python-docx's font.size only writes w:sz, so
    without this Hebrew would keep the style's inherited szCs."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), str(int(round(size_pt * 2))))


def _set_cs_bold(rPr, bold):
    """Set w:bCs (complex-script bold). Word bolds RTL/Hebrew text from
    w:bCs, not w:b - python-docx's font.bold only writes w:b."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    bCs = rPr.find(qn("w:bCs"))
    if bCs is None:
        bCs = OxmlElement("w:bCs")
        rPr.append(bCs)
    bCs.set(qn("w:val"), "1" if bold else "0")


def set_run_cs_font(run, name):
    """Set ascii/hAnsi/eastAsia/cs (complex-script) font all to `name`.

    Hebrew is rendered using the w:cs font, not w:ascii/w:hAnsi - python-docx's
    Font.name only sets ascii/hAnsi, so without this Hebrew text would silently
    fall back to Word's default complex-script font.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    _clean_rfonts(rFonts, name)


def set_run_rtl(run):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rPr = run._r.get_or_add_rPr()
    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)
    rtl.set(qn("w:val"), "1")


def set_paragraph_bidi(paragraph):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)


def rtl_paragraph(paragraph, align_right=True):
    """Make a paragraph RTL: bidi + the locked JC_MODE alignment convention.

    (align_right is kept for call-site compatibility; the actual w:jc value is
    governed by JC_MODE - see the comment at the top of this file.)
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    set_paragraph_bidi(paragraph)
    if JC_MODE == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif JC_MODE == "left":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:  # "omit" - no direct w:jc; bidi start edge (visual right) applies
        paragraph.alignment = None
    return paragraph


def set_paragraph_ltr(paragraph):
    """Force a paragraph LTR with <w:bidi w:val="0"/>. Needed because Normal
    is bidi at STYLE level (in HE docs), so Header/Footer paragraphs (based on
    Normal) inherit RTL unless explicitly overridden."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "0")


def clear_inherited_tabs(paragraph, style):
    """Emit <w:tab w:val="clear"/> entries for every tab stop the paragraph's
    style defines, so only the paragraph's own tab stops remain effective.
    Style-level and paragraph-level tab stops MERGE in Word - without the
    clears, the Footer style's default center/right stops would hijack the
    tab jumps before our stops are reached."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    style_pPr = style.element.find(qn("w:pPr"))
    if style_pPr is None:
        return
    style_tabs = style_pPr.find(qn("w:tabs"))
    if style_tabs is None:
        return
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    for st_tab in style_tabs.findall(qn("w:tab")):
        clear = OxmlElement("w:tab")
        clear.set(qn("w:val"), "clear")
        clear.set(qn("w:pos"), st_tab.get(qn("w:pos")))
        # clear entries must precede set entries; prepend
        tabs.insert(0, clear)


def add_rtl_run(paragraph, text, font=ASSISTANT, size=None, bold=None, color=None):
    """Add a run with Hebrew-correct RTL + complex-script font/size/bold set."""
    run = paragraph.add_run(text)
    set_run_rtl(run)
    set_run_cs_font(run, font)
    rPr = run._r.get_or_add_rPr()
    if size is not None:
        from docx.shared import Pt

        run.font.size = Pt(size)
        _set_cs_size(rPr, size)  # RTL text sizes from w:szCs
    if bold is not None:
        run.font.bold = bold
        _set_cs_bold(rPr, bold)  # RTL text bolds from w:bCs
    if color is not None:
        run.font.color.rgb = _rgb(color)
    return run


def set_paragraph_border(paragraph, edge, color=SAGE, sz_eighths_pt=6, space=4):
    """Add a single-line border to one edge ('top' or 'bottom') of a paragraph."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(sz_eighths_pt))
    border.set(qn("w:space"), str(space))
    border.set(qn("w:color"), color)
    pBdr.append(border)


def add_page_field(paragraph, rtl=True, font=ASSISTANT, size=None, color=None):
    """Insert a live PAGE field (not literal text) as a run sequence.

    rtl=False keeps the field runs plain LTR (for LTR footer lines).
    font/size/color apply to every run so the rendered digit is styled."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    def _style_run(run):
        if rtl:
            set_run_rtl(run)
        set_run_cs_font(run, font)
        rPr = run._r.get_or_add_rPr()
        if size is not None:
            run.font.size = Pt(size)
            _set_cs_size(rPr, size)
        if color is not None:
            run.font.color.rgb = _rgb(color)
        return run

    # begin
    r1 = _style_run(paragraph.add_run())
    fldChar_begin = OxmlElement("w:fldChar")
    fldChar_begin.set(qn("w:fldCharType"), "begin")
    r1._r.append(fldChar_begin)

    # instruction text
    r2 = _style_run(paragraph.add_run())
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    r2._r.append(instrText)

    # separate
    r3 = _style_run(paragraph.add_run())
    fldChar_sep = OxmlElement("w:fldChar")
    fldChar_sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(fldChar_sep)

    # cached display value (Word recalculates on open/print)
    r4 = _style_run(paragraph.add_run("1"))

    # end
    r5 = _style_run(paragraph.add_run())
    fldChar_end = OxmlElement("w:fldChar")
    fldChar_end.set(qn("w:fldCharType"), "end")
    r5._r.append(fldChar_end)


def style_set(style, *, font_name=None, cs_font_name=None, size=None, bold=None,
              color=None, bidi=False, align_right=False,
              space_before=None, space_after=None, line_spacing=None):
    """Modify a built-in Word style in place (font/paragraph props + Hebrew bits).

    NOTE: when cs_font_name is given, this ALSO marks the style's rPr with a
    bare <w:rtl/> (see the block below) REGARDLESS of the `bidi` flag - that
    is correct for every existing Hebrew call site (bidi is always True
    wherever cs_font_name is set) but would be wrong if reused directly for
    English styles. English styles go through set_style()/
    set_style_cs_font_only() instead, which sets the literal 4-slot fonts
    WITHOUT this side effect. Left unchanged here on purpose, to keep the
    Hebrew build path byte-for-byte identical to the approved proof."""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if font_name is not None:
        style.font.name = font_name
    if size is not None:
        style.font.size = Pt(size)
        _set_cs_size(style.element.get_or_add_rPr(), size)  # Hebrew sizes from szCs
    if bold is not None:
        style.font.bold = bold
        _set_cs_bold(style.element.get_or_add_rPr(), bold)  # Hebrew bolds from bCs
    if color is not None:
        style.font.color.rgb = _rgb(color)

    if cs_font_name is not None:
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        # sets all four literal slots AND strips theme attrs (asciiTheme etc.),
        # which otherwise override the literal names back to Calibri/TNR
        _clean_rfonts(rFonts, cs_font_name)
        # also mark the style's rPr as complex-script/rtl so it's Hebrew by default
        rtl = rPr.find(qn("w:rtl"))
        if rtl is None:
            rtl = OxmlElement("w:rtl")
            rPr.append(rtl)

    if hasattr(style, "paragraph_format") and style.paragraph_format is not None:
        pf = style.paragraph_format
        if space_before is not None:
            pf.space_before = Pt(space_before)
        if space_after is not None:
            pf.space_after = Pt(space_after)
        if line_spacing is not None:
            pf.line_spacing = line_spacing
        if align_right:
            if JC_MODE == "right":
                pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif JC_MODE == "left":
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:  # "omit"
                pf.alignment = None

    if bidi:
        pPr = style.element.get_or_add_pPr()
        bidi_el = pPr.find(qn("w:bidi"))
        if bidi_el is None:
            bidi_el = OxmlElement("w:bidi")
            pPr.append(bidi_el)


def set_style_cs_font_only(style, font_name):
    """Set the literal 4-slot rFonts on a style's rPr WITHOUT style_set()'s
    cs_font_name side effect of adding a style-level <w:rtl/> marker. Used for
    English styles: they still need literal ascii/hAnsi/eastAsia/cs fonts (the
    'every rFonts, no theme attrs' convention applies regardless of language),
    but must not be flagged complex-script/RTL by default (Caption is applied
    directly to real English paragraphs, so a stray style-level rtl marker
    there would be a real - if usually invisible - inconsistency, not just a
    cosmetic one)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    _clean_rfonts(rFonts, font_name)


def set_style(style, lang, *, font_name, size, bold, color, **kw):
    """Language-dispatching style_set(): identical to calling
    style_set(style, font_name=font_name, cs_font_name=font_name, ...) for
    lang == "he" (byte-for-byte the original call shape), and for "en" sets
    the same font/size/bold/color WITHOUT style_set's cs_font_name rtl side
    effect (via set_style_cs_font_only)."""
    if lang == "he":
        style_set(style, font_name=font_name, cs_font_name=font_name, size=size,
                  bold=bold, color=color, **kw)
    else:
        style_set(style, font_name=font_name, size=size, bold=bold, color=color, **kw)
        set_style_cs_font_only(style, font_name)


def fix_title_border(style):
    """DEFECT-1 fix: Word's built-in Title style ships with a blue (accent1,
    #4F81BD) bottom border. Recolor it to sage, 0.75pt (sz=6, eighths of a
    point), and drop the theme attributes so the literal color wins."""
    from docx.oxml.ns import qn

    pPr = style.element.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        return
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        return
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), SAGE)
    for attr in ("w:themeColor", "w:themeTint", "w:themeShade"):
        if bottom.get(qn(attr)) is not None:
            del bottom.attrib[qn(attr)]


def strip_italic(style):
    """DEFECT-2 fix: Word's built-in Subtitle style ships italic (<w:i/> +
    <w:iCs/>). Spec says Rubik REGULAR - remove both."""
    from docx.oxml.ns import qn

    rPr = style.element.get_or_add_rPr()
    for tag in ("w:i", "w:iCs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)


def fix_doc_defaults(styles_element, font_name=ASSISTANT):
    """Round-2 root cause #1 (docDefaults leg): the default template's
    rPrDefault carries ONLY theme font attributes (minorHAnsi / minorBidi),
    which resolve to Calibri (latin) and Arial (complex script). Any run that
    falls through to the defaults - e.g. one missing an explicit w:cs - lands
    on those. Point the document defaults at literal `font_name` on all four
    slots and strip the theme attrs. font_name defaults to Assistant (the
    original Hebrew body font); English builds pass Work Sans."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    docDefaults = styles_element.find(qn("w:docDefaults"))
    if docDefaults is None:
        return
    rPrDefault = docDefaults.find(qn("w:rPrDefault"))
    if rPrDefault is None:
        return
    rPr = rPrDefault.find(qn("w:rPr"))
    if rPr is None:
        return
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    _clean_rfonts(rFonts, font_name)


def fix_theme(doc, minor_font=ASSISTANT):
    """Round-2 belt-and-braces: rewrite word/theme/theme1.xml so the theme
    itself is on-brand - majorFont (headings) = Rubik always, minorFont
    (body) = `minor_font` (Assistant for Hebrew builds, Work Sans for
    English), for both the latin and cs (complex script) slots. Then even a
    style we missed that still resolves through theme attributes lands on
    Rubik/<body font> instead of Calibri/Times New Roman."""
    import re

    for part in doc.part.package.iter_parts():
        if not str(part.partname).endswith("theme1.xml"):
            continue
        xml = part.blob.decode("utf-8")

        def _patch_section(section_re, name, xml):
            m = re.search(section_re, xml, re.S)
            if not m:
                return xml
            sec = m.group(0)
            sec = re.sub(r'(<a:latin typeface=")[^"]*(")', r"\g<1>%s\g<2>" % name, sec, count=1)
            sec = re.sub(r'(<a:cs typeface=")[^"]*(")', r"\g<1>%s\g<2>" % name, sec, count=1)
            return xml[: m.start()] + sec + xml[m.end():]

        xml = _patch_section(r"<a:majorFont>.*?</a:majorFont>", RUBIK, xml)
        xml = _patch_section(r"<a:minorFont>.*?</a:minorFont>", minor_font, xml)
        part._blob = xml.encode("utf-8")
        print(f"[theme] theme1.xml majorFont->Rubik, minorFont->{minor_font}")


def add_ltr_run(paragraph, text, font=ASSISTANT, size=None, bold=None, color=None,
                no_break_hyphens=False):
    """Add a left-to-right run. Originally written for LTR islands (phone
    numbers, email addresses, URLs) inside an RTL paragraph; also reused
    as-is for ALL run content in English documents (see add_text()) since it
    already does exactly what an English run needs: literal 4-slot fonts,
    explicit <w:rtl w:val="0"/>, and szCs/bCs alongside sz/b. Deliberately
    does NOT set <w:rtl/> to "1", so the bidi algorithm treats the segment as
    an LTR island and digits/latin render in reading order (054-5244339, not
    5244339-054).

    no_break_hyphens=True replaces literal '-' with Word's native
    <w:noBreakHyphen/> so the line can never wrap mid-phone or mid-URL
    (Word treats a plain hyphen as a line-break opportunity)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if no_break_hyphens and "-" in text:
        run = paragraph.add_run()
        for i, part in enumerate(text.split("-")):
            if i:
                run._r.append(OxmlElement("w:noBreakHyphen"))
            if part:
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = part
                run._r.append(t)
    else:
        run = paragraph.add_run(text)
    set_run_cs_font(run, font)
    rPr = run._r.get_or_add_rPr()
    # explicit <w:rtl w:val="0"/> - Normal's style rPr carries <w:rtl/> in HE
    # docs, so without this the run would inherit the complex-script/RTL flag
    rtl_el = rPr.find(qn("w:rtl"))
    if rtl_el is None:
        rtl_el = OxmlElement("w:rtl")
        rPr.append(rtl_el)
    rtl_el.set(qn("w:val"), "0")
    if size is not None:
        from docx.shared import Pt

        run.font.size = Pt(size)
        _set_cs_size(rPr, size)
    if bold is not None:
        run.font.bold = bold
        _set_cs_bold(rPr, bold)
    if color is not None:
        run.font.color.rgb = _rgb(color)
    return run


def mark_table_bidi(table):
    """Mark a table <w:bidiVisual/> so column 1 renders rightmost (RTL mirror).
    <w:bidiVisual/> must come right after <w:tblStyle/> in the tblPr sequence.
    Used for every HE table; never called for EN tables (plain LTR order)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tblPr = table._tbl.tblPr
    bidi_visual = OxmlElement("w:bidiVisual")
    tblStyle_el = tblPr.find(qn("w:tblStyle"))
    if tblStyle_el is not None:
        tblStyle_el.addnext(bidi_visual)
    else:
        tblPr.insert(0, bidi_visual)


def shade_cell(cell, hex_fill):
    """Flat background fill on a table cell (used for sage-tint header rows)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


# ---------------------------------------------------------------------------
# Language dispatch layer
# ---------------------------------------------------------------------------

def body_font(lang):
    return ASSISTANT if lang == "he" else WORK_SANS


def add_text(paragraph, text, lang, font=None, size=None, bold=None, color=None):
    """Route text-run creation by document language. NOT for content that is
    always-LTR regardless of document language (phone/email/site, page-number
    digits) - those call add_ltr_run/add_page_field directly, same as the
    original script."""
    if lang == "he":
        return add_rtl_run(paragraph, text, font=font or ASSISTANT, size=size, bold=bold, color=color)
    return add_ltr_run(paragraph, text, font=font or WORK_SANS, size=size, bold=bold, color=color)


def start_paragraph(container, lang, style=None):
    """New paragraph on `container` (a Document or a table Cell), bidi-flagged
    (and start-aligned per JC_MODE) only when lang == 'he'."""
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    if lang == "he":
        rtl_paragraph(p, align_right=True)
    return p


def style_heading_runs(paragraph, lang, font=RUBIK):
    """Force cs-font (+ rtl when he) on every run of a paragraph created via
    doc.add_paragraph(text, style=...) - i.e. whose runs were NOT created
    through add_text/add_rtl_run/add_ltr_run. Mirrors the original script's
    'for r in p_title.runs: set_run_rtl(r); set_run_cs_font(r, RUBIK)' belt-
    and-suspenders pattern, applied uniformly to every heading paragraph."""
    for r in paragraph.runs:
        if lang == "he":
            set_run_rtl(r)
        set_run_cs_font(r, font)


def add_thin_rule(doc, lang):
    """A hairline sage rule (paragraph bottom border on a near-invisible 2pt-
    tall empty paragraph). Same construction as the header's trailing rule
    paragraph, factored out for reuse in the Memo's meta-block separator."""
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p = start_paragraph(doc, lang)
    for r in p.runs:
        r.font.size = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        pPr.append(rPr)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "4")
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "4")
    rPr.append(szCs)
    set_paragraph_border(p, "bottom", color=SAGE, sz_eighths_pt=6, space=2)
    return p


def add_label_value_line(doc, lang, label, value, label_suffix=": "):
    """One Normal paragraph: bold label + regular value (meta lines in Memo /
    Quote - 'To: [...]' / 'אל: [...]'). Run order = logical/reading order, so
    this works unchanged in both directions - the bidi algorithm places the
    label at the reading-start edge either way."""
    p = start_paragraph(doc, lang, style="Normal")
    font = body_font(lang)
    add_text(p, label + label_suffix, lang, font=font, size=11, bold=True, color=INK)
    add_text(p, value, lang, font=font, size=11, bold=False, color=INK)
    return p


# ---------------------------------------------------------------------------
# Content (copy) - HE strings marked "verbatim" are user-supplied; EN
# strings marked "FLAGGED" are literal translations pending Nitzan's sign-off
# (see the build report) - do not silently edit either without approval.
# ---------------------------------------------------------------------------

# header tagline - user-supplied string, verbatim (hyphen after בחינם,
# comma after שלכם, no trailing period)
TAGLINE_HEADER_HE = "מים בחינם - המותג שלכם, בידיים של כולם"
# User-approved copy (2026-07-12) - final, not a placeholder translation of
# the HE tagline. Verbatim: sentence case, single space after each period.
# This is the ONLY place this string is defined - build_header() and
# measure_pdf()'s l3 verbatim check both read TAGLINE_HEADER_EN from here,
# so there is nothing else to update in lockstep.
TAGLINE_HEADER_EN = "Hydrate for free. Elevate your brand."

WORDMARK_HE = "מרווה"
WORDMARK_EN = "MARVA"

CONTACT_PHONE = "054-5244339"
CONTACT_EMAIL = "nitzanweizmann1@gmail.com"
CONTACT_SITE = "marva-website.pages.dev"

TEXT_W_CM = 21.0 - 2.2 - 2.2  # 16.6cm text width (shared page geometry)

BODY_PARA_1_HE = (
    "זהו טקסט מציין מקום המדגים פסקת גוף רגילה במסמך. הטקסט אמור לשקף אורך וטון של "
    "תוכן אמיתי, כדי לבדוק את הריווח, גודל הגופן וקריאות השורות בעברית."
)
BODY_PARA_2_HE = (
    "פסקה נוספת לבדיקת כותרת משנה (Heading 2) והמרווח שלה מהטקסט שמעליה ומתחתיה, "
    "כולל בדיקת יישור לימין וכיווניות RTL מלאה."
)
BODY_PARA_1_EN = (
    "This is placeholder text demonstrating a normal body paragraph in the document. "
    "The text is meant to reflect the length and tone of real content, to test "
    "spacing, font size, and line readability."
)
BODY_PARA_2_EN = (
    "An additional paragraph to test a subheading (Heading 2) and its spacing from "
    "the text above and below it, including alignment and full left-to-right layout."
)

LETTER_BODY_1_HE = (
    "זהו טקסט מציין מקום לפסקת הגוף הראשונה של המכתב. יש להחליף אותו בתוכן האמיתי "
    "בעת השימוש בתבנית."
)
LETTER_BODY_2_HE = (
    "פסקה שנייה לבדיקת המרווח בין פסקאות וסגנון הגוף (Assistant, 11pt) לאורך המכתב."
)
LETTER_BODY_1_EN = (
    "This is placeholder text for the letter's first body paragraph. Replace it with "
    "real content when using this template."
)
LETTER_BODY_2_EN = (
    "A second paragraph to test spacing between paragraphs and the body style "
    "(Work Sans, 11pt) throughout the letter."
)

# ---------------------------------------------------------------------------
# COPY: per (doc_type, lang) placeholder content
# ---------------------------------------------------------------------------

COPY = {
    "report": {
        "he": dict(
            title="[כותרת הדוח]",
            subtitle="[כותרת משנה / שם הלקוח]",
            date="[תאריך]",
            h1="[כותרת ראשית]",
            body1=BODY_PARA_1_HE,
            h2="[כותרת משנית]",
            body2=BODY_PARA_2_HE,
            bullets=["[פריט ראשון]", "[פריט שני]", "[פריט שלישי]"],
            table_headers=["[עמודה א׳]", "[עמודה ב׳]", "[עמודה ג׳]"],
            table_rows=[["[נתון]", "[נתון]", "[נתון]"], ["[נתון]", "[נתון]", "[נתון]"]],
        ),
        "en": dict(
            title="[Report Title]",
            subtitle="[Subtitle / Client Name]",
            date="[Date]",
            h1="[Main Heading]",
            body1=BODY_PARA_1_EN,
            h2="[Secondary Heading]",
            body2=BODY_PARA_2_EN,
            bullets=["[First item]", "[Second item]", "[Third item]"],
            table_headers=["[Column A]", "[Column B]", "[Column C]"],
            table_rows=[["[Value]", "[Value]", "[Value]"], ["[Value]", "[Value]", "[Value]"]],
        ),
    },
    "letter": {
        "he": dict(
            date="[תאריך]",
            recipient=["[שם]", "[חברה]", "[כתובת]"],
            subject="הנדון: [נושא]",
            body1=LETTER_BODY_1_HE,
            body2=LETTER_BODY_2_HE,
            signoff=["בברכה,", "[שם]", "[תפקיד]"],
        ),
        "en": dict(
            date="[Date]",
            recipient=["[Name]", "[Company]", "[Address]"],
            subject="Re: [Subject]",
            body1=LETTER_BODY_1_EN,
            body2=LETTER_BODY_2_EN,
            signoff=["Sincerely,", "[Name]", "[Title]"],
        ),
    },
    "memo": {
        "he": dict(
            h1="תזכיר",
            meta=[("אל", "[...]"), ("מאת", "[...]"), ("תאריך", "[...]"), ("נושא", "[...]")],
            h2="[כותרת משנית]",
            body=BODY_PARA_1_HE,
            bullets=["[פריט ראשון]", "[פריט שני]", "[פריט שלישי]"],
        ),
        "en": dict(
            h1="MEMO",
            meta=[("To", "[...]"), ("From", "[...]"), ("Date", "[...]"), ("Subject", "[...]")],
            h2="[Secondary Heading]",
            body=BODY_PARA_1_EN,
            bullets=["[First item]", "[Second item]", "[Third item]"],
        ),
    },
    "quote": {
        "he": dict(
            h1="הצעת מחיר",
            meta=[("מס' הצעה", "[...]"), ("תאריך", "[...]"), ("לקוח", "[...]"), ("תוקף", "[...]")],
            table_headers=["[פריט]", "[תיאור]", "[כמות]", "[מחיר ליח']", "[סה\"כ]"],
            table_rows=[["[...]"] * 5 for _ in range(3)],
            totals=[("סה\"כ ביניים", "[...]"), ("מע\"מ 18%", "[...]"), ("סה\"כ לתשלום", "[...]")],
            payment_terms="[תנאי תשלום]",
            disclaimer="מסמך זה אינו חשבונית מס",
        ),
        "en": dict(
            h1="QUOTE",
            meta=[("Quote #", "[...]"), ("Date", "[...]"), ("Client", "[...]"), ("Valid until", "[...]")],
            table_headers=["[Item]", "[Description]", "[Qty]", "[Unit Price]", "[Total]"],
            table_rows=[["[...]"] * 5 for _ in range(3)],
            totals=[("Subtotal", "[...]"), ("VAT 18%", "[...]"), ("Total due", "[...]")],
            payment_terms="[Payment terms]",
            disclaimer="This document is not a tax invoice.",
        ),
    },
}


# ---------------------------------------------------------------------------
# Style application (parametrized version of the original inline block)
# ---------------------------------------------------------------------------

def apply_base_styles(doc, lang):
    """Configure the 5 built-in style slots (+ their linked Char styles +
    Caption) for `lang`. For lang == "he" this produces byte-identical
    styles.xml to the original Hebrew-only script (same calls, same order,
    same values). For lang == "en": bidi=False/align_right=False throughout
    (no w:bidi, default/left alignment), body font is Work Sans, and Char/
    Caption styles get their literal fonts via set_style_cs_font_only so they
    don't pick up a spurious style-level <w:rtl/> (see that function's
    docstring)."""
    styles = doc.styles
    bidi = (lang == "he")
    align = (lang == "he")
    bfont = body_font(lang)

    set_style(styles["Title"], lang, font_name=RUBIK, size=28, bold=True, color=INK,
              bidi=bidi, align_right=align)
    fix_title_border(styles["Title"])  # defect 1: blue accent1 rule -> sage 0.75pt

    set_style(styles["Subtitle"], lang, font_name=RUBIK, size=14, bold=False, color=MUTED,
              bidi=bidi, align_right=align)
    strip_italic(styles["Subtitle"])  # defect 2: drop Word's default italic
    strip_italic(styles["Subtitle Char"])  # linked char style carries it too

    set_style(styles["Heading 1"], lang, font_name=RUBIK, size=16, bold=True, color=INK,
              bidi=bidi, align_right=align, space_before=18, space_after=6)

    set_style(styles["Heading 2"], lang, font_name=RUBIK, size=13, bold=True, color=INK,
              bidi=bidi, align_right=align, space_before=12, space_after=4)

    set_style(styles["Normal"], lang, font_name=bfont, size=11, bold=False, color=INK,
              bidi=bidi, align_right=align, line_spacing=1.15, space_after=6)

    # --- purge Word's default accent1 blue (#4F81BD) from styles.xml --------
    # The linked character styles of the styles we use (and Caption, which the
    # spec defines as body-font 8.5pt muted) get on-spec run properties - the
    # color setter drops the old blue + themeColor. Everything else that still
    # carries 4F81BD is an unused latent style (Headings 3-9 variants, Intense
    # Quote/Emphasis, table color variants): delete them. Word and Google Docs
    # recreate built-ins from their own defaults if a user ever applies one.
    set_style(styles["Heading 1 Char"], lang, font_name=RUBIK, size=16, bold=True, color=INK)
    set_style(styles["Heading 2 Char"], lang, font_name=RUBIK, size=13, bold=True, color=INK)
    set_style(styles["Title Char"], lang, font_name=RUBIK, size=28, bold=True, color=INK)
    set_style(styles["Subtitle Char"], lang, font_name=RUBIK, size=14, bold=False, color=MUTED)
    set_style(styles["Caption"], lang, font_name=bfont, size=8.5, bold=False, color=MUTED)

    for st in list(styles):
        if "4F81BD" in st.element.xml:
            st.delete()

    # round-2: kill theme-font fallbacks at both remaining levels
    fix_doc_defaults(styles.element, font_name=bfont)  # docDefaults -> literal x4 slots
    fix_theme(doc, minor_font=bfont)                   # theme1.xml major=Rubik / minor=bfont


# ---------------------------------------------------------------------------
# Header / footer (shared by every doc type that shows one - all except the
# Report's title page)
# ---------------------------------------------------------------------------

def build_header(section, lang):
    """Borderless 2-cell bidiVisual(he-only) table spanning the text width.
    Cell 1 = mark + wordmark at reading-start; Cell 2 = tagline at reading-
    end. HE: bidiVisual reverses cell render order (cell1 -> visual RIGHT).
    EN: no bidiVisual, cell1 stays visual LEFT - same cell-content assignment
    either way, only the mirroring flag changes. Sage rule below, on the
    tiny trailing header paragraph."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt

    header = section.header
    header.is_linked_to_previous = False
    p_header_tail = header.paragraphs[0]

    htbl = header.add_table(rows=1, cols=2, width=Cm(TEXT_W_CM))
    p_header_tail._p.addprevious(htbl._tbl)  # table above the rule paragraph

    if lang == "he":
        mark_table_bidi(htbl)

    cell_brand = htbl.rows[0].cells[0]   # reading-start (HE: visual RIGHT via bidiVisual; EN: visual LEFT)
    cell_tag = htbl.rows[0].cells[1]     # reading-end
    cell_brand.width = Cm(8.6)
    cell_tag.width = Cm(8.0)
    cell_brand.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    cell_tag.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    p_brand = cell_brand.paragraphs[0]
    if lang == "he":
        rtl_paragraph(p_brand, align_right=True)
    p_brand.paragraph_format.space_after = Pt(0)
    run_h_img = p_brand.add_run()
    run_h_img.add_picture(PNG_MARK, height=Cm(1.0))
    run_h_space = p_brand.add_run(" ")
    if lang == "he":
        set_run_rtl(run_h_space)
    wordmark_text = WORDMARK_HE if lang == "he" else WORDMARK_EN
    add_text(p_brand, wordmark_text, lang, font=RUBIK, size=14, bold=True, color=INK)

    p_tag = cell_tag.paragraphs[0]
    tagline_text = TAGLINE_HEADER_HE if lang == "he" else TAGLINE_HEADER_EN
    if lang == "he":
        set_paragraph_bidi(p_tag)
        p_tag.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # logical right = visual LEFT under bidi
    else:
        p_tag.alignment = WD_ALIGN_PARAGRAPH.RIGHT  # plain right = reading-end for LTR
    p_tag.paragraph_format.space_after = Pt(0)
    add_text(p_tag, tagline_text, lang, font=body_font(lang), size=9, color=MUTED)

    # thin sage rule below the header content (paragraph bottom border on the
    # tiny trailing paragraph)
    if lang == "he":
        rtl_paragraph(p_header_tail, align_right=True)
    for r in p_header_tail.runs:
        r.font.size = Pt(2)
    p_header_tail.paragraph_format.space_before = Pt(0)
    p_header_tail.paragraph_format.space_after = Pt(0)
    tail_pPr = p_header_tail._p.get_or_add_pPr()
    tail_rPr = tail_pPr.find(qn("w:rPr"))
    if tail_rPr is None:
        tail_rPr = OxmlElement("w:rPr")
        tail_pPr.append(tail_rPr)  # w:rPr is the last child in the pPr sequence
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "4")  # 2pt - keeps the empty rule line near-invisible
    tail_rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "4")
    tail_rPr.append(szCs)
    set_paragraph_border(p_header_tail, "bottom", color=SAGE, sz_eighths_pt=6, space=2)


def build_footer(section, styles, lang):
    """Two lines, no tagline, no separators. Line 1 - plain LTR paragraph
    (all content is digits/Latin): phone flush LEFT, email CENTERED via a
    center tab at text-width/2, site flush RIGHT via a right tab at text-
    width. Physically IDENTICAL in both languages (contacts are already
    Latin) - only the paragraph-level bidi override differs (needed in HE,
    a no-op that's simply skipped in EN since Normal isn't bidi there).
    Line 2 - centered live PAGE field."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt

    footer = section.footer
    footer.is_linked_to_previous = False
    fkw = dict(font=body_font(lang), size=8.5, color=MUTED)

    p_f1 = footer.paragraphs[0]
    if lang == "he":
        set_paragraph_ltr(p_f1)
    p_f1.alignment = None
    clear_inherited_tabs(p_f1, styles["Footer"])
    pf1 = p_f1.paragraph_format
    pf1.tab_stops.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
    pf1.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    pf1.space_before = Pt(0)
    pf1.space_after = Pt(3)  # ~3pt between line 1 and line 2
    # sage top rule; space=6 puts ~6pt of air between the rule and line 1
    set_paragraph_border(p_f1, "top", color=SAGE, sz_eighths_pt=6, space=6)

    add_ltr_run(p_f1, CONTACT_PHONE, no_break_hyphens=True, **fkw)
    add_ltr_run(p_f1, "\t", **fkw)
    add_ltr_run(p_f1, CONTACT_EMAIL, **fkw)
    add_ltr_run(p_f1, "\t", **fkw)
    add_ltr_run(p_f1, CONTACT_SITE, no_break_hyphens=True, **fkw)

    p_f2 = footer.add_paragraph()
    if lang == "he":
        set_paragraph_ltr(p_f2)
    p_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f2.paragraph_format.space_before = Pt(0)
    p_f2.paragraph_format.space_after = Pt(0)
    add_page_field(p_f2, rtl=False, font=body_font(lang), size=8.5, color=MUTED)


# ===========================================================================
# Content builders, one per doc_type. Each receives a fresh Document() with
# apply_base_styles() already applied, and is responsible for its own
# section setup (different_first_page_header_footer), header/footer, and body.
# ===========================================================================

def build_report(doc, lang):
    """Clean title page (mark, Title, Subtitle, date - no header/footer) +
    page break + content page (header/footer + H1/body/H2/body/bullets/table).
    This is the approved proof structure - the only doc type that splits
    first-page vs. rest."""
    from docx.shared import Cm, Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT

    c = COPY["report"][lang]
    section = doc.sections[0]
    section.different_first_page_header_footer = True  # title page: no header/footer

    # ===================================================================
    # TITLE PAGE
    # ===================================================================
    p_mark = start_paragraph(doc, lang)
    p_mark.paragraph_format.space_after = Pt(12)
    run_mark = p_mark.add_run()
    run_mark.add_picture(PNG_MARK, height=Cm(3.0))

    p_title = doc.add_paragraph(c["title"], style="Title")
    if lang == "he":
        rtl_paragraph(p_title, align_right=True)
    style_heading_runs(p_title, lang, font=RUBIK)

    p_subtitle = doc.add_paragraph(c["subtitle"], style="Subtitle")
    if lang == "he":
        rtl_paragraph(p_subtitle, align_right=True)
    style_heading_runs(p_subtitle, lang, font=RUBIK)

    p_date = start_paragraph(doc, lang)
    add_text(p_date, c["date"], lang, font=body_font(lang), size=10.5, color=MUTED)

    doc.add_page_break()

    # ===================================================================
    # HEADER / FOOTER (pages 2+ only, via different_first_page)
    # ===================================================================
    build_header(section, lang)
    build_footer(section, doc.styles, lang)

    # ===================================================================
    # CONTENT PAGE
    # ===================================================================
    h1 = doc.add_paragraph(c["h1"], style="Heading 1")
    if lang == "he":
        rtl_paragraph(h1, align_right=True)
    style_heading_runs(h1, lang, font=RUBIK)

    b1 = start_paragraph(doc, lang, style="Normal")
    add_text(b1, c["body1"], lang, font=body_font(lang), size=11, color=INK)

    h2 = doc.add_paragraph(c["h2"], style="Heading 2")
    if lang == "he":
        rtl_paragraph(h2, align_right=True)
    style_heading_runs(h2, lang, font=RUBIK)

    b2 = start_paragraph(doc, lang, style="Normal")
    add_text(b2, c["body2"], lang, font=body_font(lang), size=11, color=INK)

    # bullet list - "List Bullet" built-in style for the glyph/indent, but
    # Hebrew/English font+direction is forced directly on every run (belt &
    # suspenders, since List Bullet is not one of the 5 styles we modify).
    for item in c["bullets"]:
        pb = start_paragraph(doc, lang, style="List Bullet")
        add_text(pb, item, lang, font=body_font(lang), size=11, color=INK)

    # small 3x3 table (1 header row + 2 body rows x 3 cols)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if lang == "he":
        mark_table_bidi(table)  # DEFECT-3 fix: column 1 renders rightmost

    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(c["table_headers"]):
        cell = hdr_cells[i]
        shade_cell(cell, SAGE_TINT)
        # approved-proof parity: the original script assigned .text = "" here,
        # which leaves a leading empty <w:r/> in the cell paragraph. Rendering-
        # neutral, but required for byte-stable regeneration of the approved
        # Marva-Report-HE.docx (part-by-part diff vs the approved proof showed
        # these 3 empty runs as the ONLY content difference when dropped).
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        if lang == "he":
            rtl_paragraph(p, align_right=True)
        add_text(p, htext, lang, font=RUBIK, size=10.5, bold=True, color=INK)

    for row_data in c["table_rows"]:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cell = row_cells[i]
            p = cell.paragraphs[0]
            if lang == "he":
                rtl_paragraph(p, align_right=True)
            add_text(p, val, lang, font=body_font(lang), size=10.5, color=INK)


def build_letter(doc, lang):
    """Single page, header+footer on every page (no title-page split - the
    one page IS the document). date line / blank / recipient block / blank /
    subject (Heading 2) / 2 body paragraphs / sign-off block."""
    c = COPY["letter"][lang]
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    build_header(section, lang)
    build_footer(section, doc.styles, lang)

    font = body_font(lang)

    p_date = start_paragraph(doc, lang, style="Normal")
    add_text(p_date, c["date"], lang, font=font, size=11, color=INK)

    start_paragraph(doc, lang, style="Normal")  # blank

    for line in c["recipient"]:
        p = start_paragraph(doc, lang, style="Normal")
        add_text(p, line, lang, font=font, size=11, color=INK)

    start_paragraph(doc, lang, style="Normal")  # blank

    p_subject = doc.add_paragraph(c["subject"], style="Heading 2")
    if lang == "he":
        rtl_paragraph(p_subject, align_right=True)
    style_heading_runs(p_subject, lang, font=RUBIK)

    for body in (c["body1"], c["body2"]):
        p = start_paragraph(doc, lang, style="Normal")
        add_text(p, body, lang, font=font, size=11, color=INK)

    for line in c["signoff"]:
        p = start_paragraph(doc, lang, style="Normal")
        add_text(p, line, lang, font=font, size=11, color=INK)


def build_memo(doc, lang):
    """Single page, header+footer on. 'MEMO'/'תזכיר' (Heading 1), 4 meta
    lines with bold labels, thin sage rule, then Heading 2 + body + 3
    bullets."""
    c = COPY["memo"][lang]
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    build_header(section, lang)
    build_footer(section, doc.styles, lang)

    font = body_font(lang)

    h1 = doc.add_paragraph(c["h1"], style="Heading 1")
    if lang == "he":
        rtl_paragraph(h1, align_right=True)
    style_heading_runs(h1, lang, font=RUBIK)

    for label, value in c["meta"]:
        add_label_value_line(doc, lang, label, value)

    add_thin_rule(doc, lang)

    h2 = doc.add_paragraph(c["h2"], style="Heading 2")
    if lang == "he":
        rtl_paragraph(h2, align_right=True)
    style_heading_runs(h2, lang, font=RUBIK)

    p_body = start_paragraph(doc, lang, style="Normal")
    add_text(p_body, c["body"], lang, font=font, size=11, color=INK)

    for item in c["bullets"]:
        pb = start_paragraph(doc, lang, style="List Bullet")
        add_text(pb, item, lang, font=font, size=11, color=INK)


def build_quote(doc, lang):
    """Single page, header+footer on. 'QUOTE'/'הצעת מחיר' (Heading 1), 4 meta
    lines, 5-col line-items table (sage-tint header, RTL-mirrored for he),
    narrow totals block pinned to the reading-end side, payment-terms note,
    and a caption-size 'not a tax invoice' disclaimer."""
    from docx.shared import Cm
    from docx.enum.table import WD_TABLE_ALIGNMENT

    c = COPY["quote"][lang]
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    build_header(section, lang)
    build_footer(section, doc.styles, lang)

    font = body_font(lang)

    h1 = doc.add_paragraph(c["h1"], style="Heading 1")
    if lang == "he":
        rtl_paragraph(h1, align_right=True)
    style_heading_runs(h1, lang, font=RUBIK)

    for label, value in c["meta"]:
        add_label_value_line(doc, lang, label, value)

    # ---- line-items table ---------------------------------------------
    ncols = len(c["table_headers"])
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if lang == "he":
        mark_table_bidi(table)

    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(c["table_headers"]):
        cell = hdr_cells[i]
        shade_cell(cell, SAGE_TINT)
        p = cell.paragraphs[0]
        if lang == "he":
            rtl_paragraph(p, align_right=True)
        add_text(p, htext, lang, font=RUBIK, size=10.5, bold=True, color=INK)

    for row_data in c["table_rows"]:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cell = row_cells[i]
            p = cell.paragraphs[0]
            if lang == "he":
                rtl_paragraph(p, align_right=True)
            add_text(p, val, lang, font=font, size=10.5, color=INK)

    # ---- narrow totals block -------------------------------------------
    # Pinned to the READING-END side: visual LEFT for he (RTL end, matching
    # the mirrored table's leftmost "total" column), visual RIGHT for en.
    # EMPIRICAL (r1 failure, first full run): for a <w:bidiVisual/> table
    # Word interprets tblPr w:jc LOGICALLY, exactly like paragraph w:jc under
    # w:bidi - jc="left" rendered the HE block at the visual RIGHT (totals
    # label x0=481.66 on a 595pt page). So BOTH languages want jc="right":
    # HE reads it logically (right = end-of-reading = visual LEFT) and EN
    # reads it literally (visual RIGHT = its reading end). Internal
    # bidiVisual (he only) keeps the bold label at the reading-start edge
    # within the narrow block itself.
    totals = doc.add_table(rows=1, cols=2)
    if lang == "he":
        mark_table_bidi(totals)
    totals.alignment = WD_TABLE_ALIGNMENT.RIGHT

    def _fill_totals_row(cells, label, value):
        cells[0].width = Cm(4.5)
        cells[1].width = Cm(2.5)
        p_label = cells[0].paragraphs[0]
        if lang == "he":
            rtl_paragraph(p_label, align_right=True)
        add_text(p_label, label, lang, font=font, size=11, bold=True, color=INK)
        p_value = cells[1].paragraphs[0]
        if lang == "he":
            rtl_paragraph(p_value, align_right=True)
        add_text(p_value, value, lang, font=font, size=11, color=INK)

    _fill_totals_row(totals.rows[0].cells, *c["totals"][0])
    for label, value in c["totals"][1:]:
        _fill_totals_row(totals.add_row().cells, label, value)

    p_terms = start_paragraph(doc, lang, style="Normal")
    add_text(p_terms, c["payment_terms"], lang, font=font, size=11, color=INK)

    p_disclaimer = doc.add_paragraph(style="Caption")
    if lang == "he":
        rtl_paragraph(p_disclaimer, align_right=True)
    add_text(p_disclaimer, c["disclaimer"], lang, font=font, size=8.5, color=MUTED)


BUILDERS = {
    "report": build_report,
    "letter": build_letter,
    "memo": build_memo,
    "quote": build_quote,
}

DOC_TYPES = ("report", "letter", "memo", "quote")
LANGS = ("he", "en")


def out_paths(doc_type, lang):
    stem = f"Marva-{doc_type.capitalize()}-{lang.upper()}"
    return os.path.join(OUT_DIR, stem + ".docx"), os.path.join(OUT_DIR, stem + ".pdf")


def build_doc(doc_type, lang, out_path=None):
    """Master entry point: fresh Document(), shared page geometry, base
    styles, then dispatch to the doc_type's content builder."""
    from docx import Document
    from docx.shared import Cm

    docx_path = out_path or out_paths(doc_type, lang)[0]
    os.makedirs(OUT_DIR, exist_ok=True)

    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    apply_base_styles(doc, lang)
    BUILDERS[doc_type](doc, lang)

    doc.save(docx_path)
    print(f"[docx] saved: {docx_path} (doc_type={doc_type}, lang={lang}, JC_MODE={JC_MODE})")
    return docx_path


# backward-compatible alias matching the original script's entry point
def build_docx(jc_mode=None, out_path=None):
    global JC_MODE
    if jc_mode is not None:
        JC_MODE = jc_mode
    return build_doc("report", "he", out_path=out_path or DOCX_OUT)


# ---------------------------------------------------------------------------
# Verification (docx-level, structural - language/doc-type agnostic)
# ---------------------------------------------------------------------------

def verify_docx(path):
    """Re-open the docx with python-docx and confirm bidi flags + PAGE field exist."""
    from docx import Document
    import zipfile

    assert os.path.exists(path), f"missing {path}"
    size = os.path.getsize(path)
    assert size > 0, f"{path} is empty"

    doc = Document(path)  # must open without raising

    with zipfile.ZipFile(path) as z:
        document_xml = z.read("word/document.xml").decode("utf-8")
        header_xmls = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/header")]
        footer_xmls = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("word/footer")]
        styles_xml = z.read("word/styles.xml").decode("utf-8")

    all_xml = document_xml + "".join(header_xmls) + "".join(footer_xmls)

    checks = {
        "w:rtl present (run RTL)": "<w:rtl/>" in all_xml or "<w:rtl " in all_xml,
        "PAGE field present": "PAGE" in all_xml and "fldChar" in all_xml,
        "styles.xml has Rubik": "Rubik" in styles_xml,
        "file size > 0": size > 0,
    }
    for name, ok in checks.items():
        print(f"[verify docx] {'OK ' if ok else 'FAIL'} - {name}")
    if not all(checks.values()):
        raise AssertionError(f"docx verification failed: {checks}")
    print(f"[verify docx] all checks passed ({size} bytes)")


def ensure_fonts_loaded():
    """Round-2 environment fix: per-user fonts that were installed by copying
    files + writing HKCU registry entries are NOT visible to apps until the
    next logon, because nothing called AddFontResource / broadcast
    WM_FONTCHANGE (Explorer's right-click Install does both). Word then
    silently substitutes Calibri/Arial at PDF export - even though it lists
    the names in FontNames (that list includes Office *cloud* fonts, so it
    proves nothing about local availability). Loading here is idempotent and
    session-scoped; the registry entries make the install permanent across
    logons."""
    import ctypes
    import glob

    try:
        gdi32 = ctypes.WinDLL("gdi32")
        user32 = ctypes.WinDLL("user32")
    except OSError:
        return  # not Windows
    folder = os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts")
    n = 0
    for f in glob.glob(os.path.join(folder, "*.ttf")):
        n += gdi32.AddFontResourceW(f)
    if n:
        HWND_BROADCAST, WM_FONTCHANGE, SMTO_ABORTIFHUNG = 0xFFFF, 0x001D, 0x0002
        res = ctypes.c_ulong()
        user32.SendMessageTimeoutW(HWND_BROADCAST, WM_FONTCHANGE, 0, 0,
                                   SMTO_ABORTIFHUNG, 1000, ctypes.byref(res))
        print(f"[fonts] {n} user font resources loaded into this session")


def export_pdf(docx_path=None, pdf_path=None):
    docx_path = docx_path or DOCX_OUT
    pdf_path = pdf_path or PDF_OUT
    os.makedirs(OUT_DIR, exist_ok=True)
    ensure_fonts_loaded()
    try:
        from docx2pdf import convert

        convert(docx_path, pdf_path)
    except Exception as e:
        print(f"[pdf] docx2pdf failed ({e}), falling back to raw Word COM automation")
        import win32com.client

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        try:
            wdoc = word.Documents.Open(docx_path)
            wdoc.SaveAs(pdf_path, FileFormat=17)  # wdFormatPDF
            wdoc.Close(False)
        finally:
            word.Quit()

    assert os.path.exists(pdf_path), "PDF export did not produce a file"
    size = os.path.getsize(pdf_path)
    print(f"[pdf] saved: {pdf_path} ({size} bytes)")
    assert size > 20 * 1024, f"PDF is suspiciously small ({size} bytes), expected >20KB"
    return pdf_path


# ---------------------------------------------------------------------------
# Render-level QA (ground truth for the RTL alignment convention + direction-
# and doc-type-aware structural checks)
# ---------------------------------------------------------------------------

def _contains(text, needle, lang):
    """Match `needle` in extracted PDF text. For he, also try the reversed
    string - PDF text extraction order for Hebrew is renderer-dependent
    (logical or visual), same empirical finding as the original script's
    _contains_he. English text extraction doesn't need this."""
    if needle in text:
        return True
    if lang == "he" and needle[::-1] in text:
        return True
    return False


def spans(page):
    out = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                out.append(span)
    return out


def find_span(page, needle, lang, y_max=None, y_min=None):
    """Best (widest) span whose text contains `needle` (direction-aware),
    optionally restricted to a y-band."""
    best = None
    for s in spans(page):
        if not _contains(s["text"], needle, lang):
            continue
        y0 = s["bbox"][1]
        if y_max is not None and y0 > y_max:
            continue
        if y_min is not None and y0 < y_min:
            continue
        if best is None or s["bbox"][2] > best["bbox"][2]:
            best = s
    return best


def _needle_core(t):
    """Letter/digit core of a table-header needle: strip the placeholder
    brackets AND any edge punctuation. PDF bidi extraction displaces
    punctuation at the EDGES of an RTL cell text - observed on quote-he,
    where '[מחיר ליח\\']' extracted as ']\\'[מחיר ליח' (the trailing
    apostrophe migrated to sit beside the closing bracket), so a needle that
    keeps edge punctuation can never match contiguously. INTERNAL marks are
    kept: the gershayim in סה\"כ is not at an edge and extracts in place."""
    return t.strip("[]'\"׳״ ")


def row_extremes(page, lang, col_texts):
    """Given header texts in LOGICAL (definition) order, locate the header ROW
    and return (leftmost_text, rightmost_text) by x-position among the column
    headers. Used for table column-order checks (Report's 3-col table, Quote's
    5-col table) in both directions. Pass needles already reduced via
    _needle_core() so matching and caller-side expectations share strings.

    Band-FIRST strategy (check-bug fix from the first full run): needles like
    'סה"כ'/'Total' also occur in the quote's totals block on other rows, and
    the old find-widest-span-page-wide approach latched onto 'סה"כ לתשלום' /
    'Total due' there, escaped the header band, and reported a false column-
    order failure while the rendered table was correct. Now: anchor on a col
    text that matches exactly ONE span page-wide, take its y-band as the
    header row, and match every col text against spans INSIDE that band only.
    Requires ALL columns found in the band (stricter than the old >=2)."""
    all_spans = spans(page)

    def hits(t):
        return [s for s in all_spans if _contains(s["text"], t, lang)]

    anchor_span = None
    for t in col_texts:
        h = hits(t)
        if len(h) == 1:
            anchor_span = h[0]
            break
    if anchor_span is None:
        return None, None
    y0, y1 = anchor_span["bbox"][1], anchor_span["bbox"][3]
    band_spans = [s for s in all_spans if s["bbox"][1] < y1 and s["bbox"][3] > y0]
    found = []
    for t in col_texts:
        in_band = [s for s in band_spans if _contains(s["text"], t, lang)]
        if in_band:
            found.append((t, max(in_band, key=lambda s: s["bbox"][2] - s["bbox"][0])))
    if len(found) < len(col_texts):
        return None, None
    found.sort(key=lambda ts: ts[1]["bbox"][0])
    return found[0][0], found[-1][0]


def measure_pdf(pdf_path, doc_type, lang):
    """Measure text bounding boxes in the rendered PDF. Returns a dict of
    metrics used to verify the RTL/LTR mirroring convention and doc-specific
    structure. `header_page` is the 0-indexed page that carries the running
    header/footer: page 2 (index 1) for report (title page is clean), page 1
    (index 0) for letter/memo/quote (single-section, header everywhere)."""
    import re as _re
    import fitz

    c = COPY[doc_type][lang]
    doc = fitz.open(pdf_path)
    m = {"page_count": doc.page_count, "page_width": doc[0].rect.width}

    header_page_idx = 1 if doc_type == "report" else 0
    hp = doc[header_page_idx]
    width = hp.rect.width
    height = hp.rect.height
    m["page_width"] = width
    right_third = width * 0.66
    left_third = width * 0.34

    def in_start_third(x1_or_x0, is_x1):
        # HE: expect the RIGHT third (use x1, the right edge); EN: expect the
        # LEFT third (use x0, the left edge). Caller passes whichever edge is
        # meaningful; this just applies the threshold for the doc's language.
        if lang == "he":
            return x1_or_x0 is not None and x1_or_x0 > right_third
        return x1_or_x0 is not None and x1_or_x0 < left_third

    # ---- header wordmark + tagline -------------------------------------
    wordmark_text = WORDMARK_HE if lang == "he" else WORDMARK_EN
    wordmark = find_span(hp, wordmark_text, lang, y_max=100)
    m["wordmark_x1"] = wordmark["bbox"][2] if wordmark else None
    m["wordmark_x0"] = wordmark["bbox"][0] if wordmark else None
    m["wordmark_edge_ok"] = in_start_third(m["wordmark_x1"] if lang == "he" else m["wordmark_x0"], lang == "he")

    tagline_text = TAGLINE_HEADER_HE if lang == "he" else TAGLINE_HEADER_EN
    anchor_word = tagline_text.split()[0]
    anchor = find_span(hp, anchor_word, lang, y_max=110)
    m["tagline_x0"] = None
    m["tagline_x1"] = None
    m["tagline_text_ok"] = None
    m["tagline_same_band"] = None
    m["tagline_extracted"] = None
    if anchor:
        ay0, ay1 = anchor["bbox"][1], anchor["bbox"][3]
        if lang == "he":
            # header band, left of the wordmark cell - collect right-to-left
            band = [s for s in spans(hp)
                    if s["bbox"][1] < ay1 and s["bbox"][3] > ay0
                    and s["bbox"][0] < width * 0.6]
            band.sort(key=lambda s: -s["bbox"][0])

            def _rtl_norm(t):
                stripped = t.lstrip()
                return stripped + " " * (len(t) - len(stripped))

            joined = "".join(_rtl_norm(s["text"]) for s in band)
        else:
            # header band, right of the wordmark cell - collect left-to-right
            band = [s for s in spans(hp)
                    if s["bbox"][1] < ay1 and s["bbox"][3] > ay0
                    and s["bbox"][0] > width * 0.4]
            band.sort(key=lambda s: s["bbox"][0])
            joined = "".join(s["text"] for s in band)

        # trailing-space spans (Word PDF export emits the cell-end space as
        # its own span) would drag x1 rightward of the real text - measure
        # the edges from non-blank spans only
        text_band = [s for s in band if s["text"].strip()]
        m["tagline_x0"] = min(s["bbox"][0] for s in text_band) if text_band else None
        m["tagline_x1"] = max(s["bbox"][2] for s in text_band) if text_band else None
        cleaned = _re.sub("[‎‏‪-‮⁦-⁩]", "", joined)
        cleaned = _re.sub(r"\s+", " ", cleaned).strip()
        target = _re.sub(r"\s+", " ", tagline_text).strip()
        m["tagline_text_ok"] = cleaned in (target, target[::-1])
        m["tagline_extracted"] = cleaned
        if wordmark:
            m["tagline_same_band"] = (ay0 < wordmark["bbox"][3] and ay1 > wordmark["bbox"][1])
    # Reading-END anchoring, direction-true (check-bug fix from the first
    # full run): the tagline is anchored at its reading-end edge, so the
    # meaningful edge is the one NEAR that end - HE: LEFT edge x0 in the left
    # third (approved-proof convention, unchanged); EN mirror: RIGHT edge x1
    # in the right third. The old EN check asserted x0 > right_third, which a
    # ~163pt-long tagline legitimately violates (measured x0=325.87,
    # x1=489.28 with right_third=392.91 - the render was correct).
    if lang == "he":
        m["tagline_edge_ok"] = m["tagline_x0"] is not None and m["tagline_x0"] < width * 0.33
    else:
        m["tagline_edge_ok"] = m["tagline_x1"] is not None and m["tagline_x1"] > width * 0.67

    # ---- footer (phone LEFT / email CENTER / site RIGHT + page number) ----
    LEFT_MARGIN = 2.2 * 72 / 2.54
    m["left_margin"] = LEFT_MARGIN
    m["right_margin_x"] = width - LEFT_MARGIN
    m["center_x"] = width / 2

    footer_spans = [s for s in spans(hp) if s["bbox"][1] > height * 0.85]
    phone_s = next((s for s in footer_spans if "054" in s["text"]), None)
    email_s = next((s for s in footer_spans if "@" in s["text"]), None)
    site_s = next((s for s in footer_spans if "pages" in s["text"]), None)
    m["phone_x0"] = phone_s["bbox"][0] if phone_s else None
    m["email_center"] = (email_s["bbox"][0] + email_s["bbox"][2]) / 2 if email_s else None
    m["site_x1"] = site_s["bbox"][2] if site_s else None
    m["footer_line1_same_y"] = None
    line1_y1 = None
    if phone_s and email_s and site_s:
        y0s = [s["bbox"][1] for s in (phone_s, email_s, site_s)]
        y1s = [s["bbox"][3] for s in (phone_s, email_s, site_s)]
        m["footer_line1_same_y"] = max(y0s) < min(y1s)
        line1_y1 = max(y1s)

    m["pagenum_center"] = None
    m["pagenum_below_line1"] = None
    if line1_y1 is not None:
        digit_spans = [s for s in footer_spans
                       if s["text"].strip().isdigit() and s["bbox"][1] >= line1_y1 - 1]
        if digit_spans:
            d = min(digit_spans, key=lambda s: abs((s["bbox"][0] + s["bbox"][2]) / 2 - width / 2))
            m["pagenum_center"] = (d["bbox"][0] + d["bbox"][2]) / 2
            m["pagenum_below_line1"] = d["bbox"][1] >= line1_y1 - 1

    footer_text = hp.get_text()
    # hyphen-minus / hyphen / no-break hyphen / minus - noBreakHyphen may
    # extract as any of these depending on the (substituted) font's cmap
    hyph = "[-‐‑−]"
    # \s* on both sides of the hyphen (check-bug fix from the first full
    # run): Word's PDF writer may chunk the phone into separate text runs -
    # observed on every EN doc as '054' / '-' / '5244339' spans with ~2pt
    # gaps, which PyMuPDF renders into the text stream as virtual spaces
    # ('054 - 5244339') while the visual line is a correct, contiguous
    # 054-5244339 at the left margin. Tolerance is applied to the positive
    # AND the scramble pattern alike, so scramble detection is not weakened.
    m["phone_reads_ltr"] = bool(_re.search(r"054\s*" + hyph + r"\s*5244339", footer_text))
    m["phone_scrambled"] = bool(_re.search(r"5244339\s*" + hyph + r"\s*054", footer_text))

    # ---- doc-type-specific structural anchors --------------------------
    if doc_type == "report":
        p1 = doc[0]
        title_needle = c["title"].strip("[]")
        date_needle = c["date"].strip("[]")
        h1_needle = c["h1"].strip("[]")
        title_span = find_span(p1, title_needle, lang)
        date_span = find_span(p1, date_needle, lang)
        h1_span = find_span(hp, h1_needle, lang)
        m["title_x1"] = title_span["bbox"][2] if title_span else None
        m["title_x0"] = title_span["bbox"][0] if title_span else None
        m["date_x1"] = date_span["bbox"][2] if date_span else None
        m["date_x0"] = date_span["bbox"][0] if date_span else None
        m["h1_x1"] = h1_span["bbox"][2] if h1_span else None
        m["h1_x0"] = h1_span["bbox"][0] if h1_span else None
        if lang == "he":
            m["title_edge_ok"] = m["title_x1"] is not None and m["title_x1"] > right_third
            m["date_edge_ok"] = m["date_x1"] is not None and m["date_x1"] > right_third
            m["h1_edge_ok"] = m["h1_x1"] is not None and m["h1_x1"] > right_third
        else:
            m["title_edge_ok"] = m["title_x0"] is not None and m["title_x0"] < left_third
            m["date_edge_ok"] = m["date_x0"] is not None and m["date_x0"] < left_third
            m["h1_edge_ok"] = m["h1_x0"] is not None and m["h1_x0"] < left_third

        # title page must be clean (no header/footer content)
        m["titlepage_wordmark_absent"] = find_span(p1, wordmark_text, lang, y_max=100) is None
        m["titlepage_phone_absent"] = "054" not in p1.get_text()

        needles = [_needle_core(t) for t in c["table_headers"]]
        left, right = row_extremes(hp, lang, needles)
        expect_first, expect_last = needles[0], needles[-1]
        if lang == "he":
            m["table_order_ok"] = (right == expect_first and left == expect_last)
        else:
            m["table_order_ok"] = (left == expect_first and right == expect_last)
        m["table_order_detail"] = f"left={left!r} right={right!r}"

    elif doc_type == "letter":
        subj_needle = "הנדון" if lang == "he" else "Subject"
        subj_span = find_span(hp, subj_needle, lang)
        m["subject_x1"] = subj_span["bbox"][2] if subj_span else None
        m["subject_x0"] = subj_span["bbox"][0] if subj_span else None
        if lang == "he":
            m["subject_edge_ok"] = m["subject_x1"] is not None and m["subject_x1"] > right_third
        else:
            m["subject_edge_ok"] = m["subject_x0"] is not None and m["subject_x0"] < left_third

    elif doc_type == "memo":
        h1_needle = c["h1"]
        h1_span = find_span(hp, h1_needle, lang, y_max=height * 0.5)
        m["h1_x1"] = h1_span["bbox"][2] if h1_span else None
        m["h1_x0"] = h1_span["bbox"][0] if h1_span else None
        if lang == "he":
            m["h1_edge_ok"] = m["h1_x1"] is not None and m["h1_x1"] > right_third
        else:
            m["h1_edge_ok"] = m["h1_x0"] is not None and m["h1_x0"] < left_third

    elif doc_type == "quote":
        h1_needle = c["h1"]
        h1_span = find_span(hp, h1_needle, lang, y_max=height * 0.5)
        m["h1_x1"] = h1_span["bbox"][2] if h1_span else None
        m["h1_x0"] = h1_span["bbox"][0] if h1_span else None
        if lang == "he":
            m["h1_edge_ok"] = m["h1_x1"] is not None and m["h1_x1"] > right_third
        else:
            m["h1_edge_ok"] = m["h1_x0"] is not None and m["h1_x0"] < left_third

        needles = [_needle_core(t) for t in c["table_headers"]]
        left, right = row_extremes(hp, lang, needles)
        expect_first, expect_last = needles[0], needles[-1]
        if lang == "he":
            m["table_order_ok"] = (right == expect_first and left == expect_last)
        else:
            m["table_order_ok"] = (left == expect_first and right == expect_last)
        m["table_order_detail"] = f"left={left!r} right={right!r}"

        subtotal_label = c["totals"][0][0]
        subtotal_span = find_span(hp, subtotal_label, lang)
        m["totals_x0"] = subtotal_span["bbox"][0] if subtotal_span else None
        if lang == "he":
            m["totals_side_ok"] = m["totals_x0"] is not None and m["totals_x0"] < width * 0.5
        else:
            m["totals_side_ok"] = m["totals_x0"] is not None and m["totals_x0"] > width * 0.5

        disc_needle = c["disclaimer"][:12]
        m["disclaimer_present"] = find_span(hp, disc_needle, lang) is not None

    doc.close()
    return m


def qa_checks(doc_type, lang, docx_path=None, pdf_path=None):
    """Post-build acceptance checks. All checks under `checks` must pass on
    every build; `known_blocked` (embedded-font ground truth) is reported but
    does not fail the build - see the i1-i3 note below."""
    import zipfile
    import re

    if docx_path is None or pdf_path is None:
        dp, pp = out_paths(doc_type, lang)
        docx_path = docx_path or dp
        pdf_path = pdf_path or pp

    c = COPY[doc_type][lang]

    with zipfile.ZipFile(docx_path) as z:
        document_xml = z.read("word/document.xml").decode("utf-8")
        styles_xml = z.read("word/styles.xml").decode("utf-8")
        header_xml = "".join(
            z.read(n).decode("utf-8") for n in z.namelist()
            if n.startswith("word/header")
        )
        footer_xml = "".join(
            z.read(n).decode("utf-8") for n in z.namelist()
            if n.startswith("word/footer")
        )
        hf_xml = header_xml + footer_xml

    title_m = re.search(r'<w:style [^>]*w:styleId="Title".*?</w:style>', styles_xml, re.S)
    subtitle_m = re.search(r'<w:style [^>]*w:styleId="Subtitle".*?</w:style>', styles_xml, re.S)
    title_xml = title_m.group(0) if title_m else ""
    subtitle_xml = subtitle_m.group(0) if subtitle_m else ""

    content_xml = document_xml + hf_xml
    applied_colors = set(re.findall(r'w:(?:val|fill|color)="([0-9A-Fa-f]{6})"', content_xml))
    allowed = {INK, MUTED, SAGE, SAGE_TINT}

    metrics = measure_pdf(pdf_path, doc_type, lang)
    expected_pages = 2 if doc_type == "report" else 1

    # bidiVisual invariant: HE = 1 (header table) + number of content tables;
    # EN = 0, always (see mark_table_bidi docstring / module docstring)
    n_content_tables = document_xml.count("<w:tbl>")
    expected_bidi = (1 + n_content_tables) if lang == "he" else 0
    actual_bidi = content_xml.count("<w:bidiVisual/>")

    # (g) no theme font attrs in the rFonts of the five modified styles or
    # docDefaults - theme attrs override literal names (round-2 root cause 1)
    theme_attr_re = re.compile(r"w:(?:ascii|hAnsi|eastAsia)Theme=|w:cstheme=")
    g_blocks = {}
    for sid in ("Title", "Subtitle", "Heading1", "Heading2", "Normal"):
        sm = re.search(r'<w:style [^>]*w:styleId="%s".*?</w:style>' % sid, styles_xml, re.S)
        g_blocks[sid] = sm.group(0) if sm else ""
    dm = re.search(r"<w:docDefaults>.*?</w:docDefaults>", styles_xml, re.S)
    g_blocks["docDefaults"] = dm.group(0) if dm else ""
    g_offenders = [name for name, block in g_blocks.items()
                   if any(theme_attr_re.search(rf) for rf in re.findall(r"<w:rFonts[^>]*/>", block))]
    g_literal_ok = all(
        all(f'w:{slot}="' in rf for slot in ("ascii", "hAnsi", "eastAsia", "cs"))
        for name, block in g_blocks.items()
        for rf in re.findall(r"<w:rFonts[^>]*/>", block)
    )

    # (h) every run-level rFonts that names w:ascii also names w:cs with the
    # same value (RTL runs render Hebrew from the cs slot - root cause 2)
    h_bad = []
    for rf in re.findall(r"<w:rFonts[^>]*/>", content_xml):
        m_ascii = re.search(r'w:ascii="([^"]*)"', rf)
        m_cs = re.search(r'w:cs="([^"]*)"', rf)
        if m_ascii and (not m_cs or m_cs.group(1) != m_ascii.group(1)):
            h_bad.append(rf)

    # (i) GROUND TRUTH - fonts actually embedded in the Word-exported PDF
    import fitz

    pdf_doc = fitz.open(pdf_path)
    embedded = set()
    for page in pdf_doc:
        for f in page.get_fonts(full=True):
            embedded.add(f[3])  # basefont name, e.g. 'AAAAAA+Rubik-Bold'
    pdf_doc.close()
    has_rubik = any("Rubik" in n for n in embedded)
    expected_body_font = "Assistant" if lang == "he" else "WorkSans"
    has_body_font = any(expected_body_font.replace(" ", "") in n.replace(" ", "") for n in embedded)
    off_brand = sorted(n for n in embedded
                       if any(bad in n for bad in ("Calibri", "TimesNewRoman", "ArialMT")))

    checks = {
        # (a) defects 1+2 (style-level, checked regardless of doc_type - all
        # 8 files share the same 5 modified built-in styles)
        "a1: no 4F81BD anywhere in styles.xml": "4F81BD" not in styles_xml,
        'a2: no themeColor="accent1" on Title': 'themeColor="accent1"' not in title_xml,
        "a3: Subtitle has no <w:i/>": "<w:i/>" not in subtitle_xml,
        "a4: Subtitle has no <w:iCs/>": "<w:iCs/>" not in subtitle_xml,
        # (b) bidiVisual invariant (mirrors HE, plain order EN)
        f"b: bidiVisual count matches mirroring convention (expected={expected_bidi}, actual={actual_bidi})":
            actual_bidi == expected_bidi,
        # (e) regressions, language-aware
        "e1: w:bidi present" if lang == "he" else "e1: no w:bidi anywhere (LTR, no bidi)":
            (("<w:bidi/>" in content_xml or "<w:bidi>" in content_xml) if lang == "he"
             else ("<w:bidi" not in document_xml)),
        "e2: w:rtl(1) present on Hebrew runs" if lang == "he" else "e2: no w:rtl val=1 in EN document body":
            ('<w:rtl w:val="1"/>' in document_xml if lang == "he"
             else ('<w:rtl w:val="1"/>' not in document_xml)),
        "e3: real PAGE field present": "PAGE" in hf_xml and "fldChar" in hf_xml,
        f"e4: only palette colors applied ({sorted(applied_colors)})": applied_colors <= allowed,
        "e5: PDF > 20KB": os.path.getsize(pdf_path) > 20 * 1024,
        # footer bidi / geometry (identical construction across all 8 files)
        "f1: phone reads LTR (054-5244339)": metrics["phone_reads_ltr"],
        "f2: phone not scrambled": not metrics["phone_scrambled"],
        # (g)/(h) font-slot correctness (language/doc-type agnostic)
        f"g1: no theme font attrs in 5 styles+docDefaults rFonts (offenders={g_offenders})": not g_offenders,
        "g2: literal ascii/hAnsi/eastAsia/cs present in those rFonts": g_literal_ok,
        f"h: every run rFonts with w:ascii has matching w:cs ({len(h_bad)} bad)": not h_bad,
        # (j) footer line 1: phone LEFT / email CENTER / site RIGHT, one line
        f"j1: phone x0 within 36pt of left margin (x0={metrics['phone_x0']}, margin={metrics['left_margin']:.1f})":
            metrics["phone_x0"] is not None and abs(metrics["phone_x0"] - metrics["left_margin"]) <= 36,
        f"j2: email center within 30pt of page center (c={metrics['email_center']}, center={metrics['center_x']:.1f})":
            metrics["email_center"] is not None and abs(metrics["email_center"] - metrics["center_x"]) <= 30,
        f"j3: site x1 within 36pt of right margin (x1={metrics['site_x1']}, margin={metrics['right_margin_x']:.1f})":
            metrics["site_x1"] is not None and abs(metrics["site_x1"] - metrics["right_margin_x"]) <= 36,
        "j4: phone/email/site on the same line": metrics["footer_line1_same_y"] is True,
        # (k) page-number line centered, below line 1
        f"k1: page number center within 30pt of page center (c={metrics['pagenum_center']})":
            metrics["pagenum_center"] is not None and abs(metrics["pagenum_center"] - metrics["center_x"]) <= 30,
        "k2: page number below line 1": metrics["pagenum_below_line1"] is True,
        # (l) header: wordmark at reading-start + tagline at reading-end, same band, verbatim
        f"l0: wordmark found and at reading-start edge (x1={metrics['wordmark_x1']}, x0={metrics['wordmark_x0']})":
            metrics["wordmark_edge_ok"],
        f"l1: tagline at reading-end edge (x0={metrics['tagline_x0']}, x1={metrics['tagline_x1']})": metrics["tagline_edge_ok"],
        "l2: tagline in same header band as wordmark": metrics["tagline_same_band"] is True,
        "l3: tagline text verbatim": metrics["tagline_text_ok"] is True,
        # (m) footer content limits
        "m1: footer has no tagline text": (TAGLINE_HEADER_HE.split()[0] if lang == "he" else TAGLINE_HEADER_EN.split()[0]) not in footer_xml,
        "m2: footer has no · separators": "·" not in footer_xml,
        # (n) page count matches doc type (report=2, else=1)
        f"n: page count == {expected_pages} (actual={metrics['page_count']})": metrics["page_count"] == expected_pages,
    }

    if doc_type == "report":
        checks.update({
            f"c1: title at start-of-reading edge (x1={metrics['title_x1']}, x0={metrics['title_x0']})": metrics["title_edge_ok"],
            f"c2: date at start-of-reading edge (x1={metrics['date_x1']}, x0={metrics['date_x0']})": metrics["date_edge_ok"],
            f"c4: Heading 1 at start-of-reading edge (x1={metrics['h1_x1']}, x0={metrics['h1_x0']})": metrics["h1_edge_ok"],
            f"d: table column order matches mirroring convention ({metrics['table_order_detail']})": metrics["table_order_ok"],
            "p1: title page has no header wordmark": metrics["titlepage_wordmark_absent"],
            "p2: title page has no footer phone": metrics["titlepage_phone_absent"],
        })
    elif doc_type == "letter":
        checks.update({
            f"c: subject (Heading 2) at start-of-reading edge (x1={metrics['subject_x1']}, x0={metrics['subject_x0']})": metrics["subject_edge_ok"],
        })
    elif doc_type == "memo":
        checks.update({
            f"c: 'MEMO'/'תזכיר' heading at start-of-reading edge (x1={metrics['h1_x1']}, x0={metrics['h1_x0']})": metrics["h1_edge_ok"],
            # python-docx serializes bold=True as an EMPTY <w:b/> (no w:val
            # attribute - True is the schema default); the first-run check
            # counted '<w:b w:val="1"/>' which python-docx never writes, and
            # failed while the label runs were verifiably bold in the XML
            # (<w:b/> + <w:bCs w:val="1"/> pairs). Count both slots: w:b for
            # Latin, w:bCs for the Hebrew/complex-script bold.
            "q: at least 4 bold label runs, w:b + w:bCs pairs (meta block)":
                document_xml.count("<w:b/>") >= 4 and document_xml.count('<w:bCs w:val="1"/>') >= 4,
        })
    elif doc_type == "quote":
        checks.update({
            f"c: 'QUOTE'/'הצעת מחיר' heading at start-of-reading edge (x1={metrics['h1_x1']}, x0={metrics['h1_x0']})": metrics["h1_edge_ok"],
            f"d: line-items table column order matches mirroring convention ({metrics['table_order_detail']})": metrics["table_order_ok"],
            f"r1: totals block on reading-end side (x0={metrics['totals_x0']})": metrics["totals_side_ok"],
            "r2: 'not a tax invoice' disclaimer present": metrics["disclaimer_present"],
            # same <w:b/> serialization note as the memo q-check; quote floor
            # is 12: 4 meta labels + 5 table header cells + 3 totals labels
            "q: at least 12 bold runs, w:b + w:bCs pairs (4 meta + 5 header + 3 totals)":
                document_xml.count("<w:b/>") >= 12 and document_xml.count('<w:bCs w:val="1"/>') >= 12,
        })

    # (i) ground truth - embedded fonts in the exported PDF. KNOWN-BLOCKED
    # until the user signs out/in (per-user font install not visible to Word
    # in this session) - reported, but does not fail the build.
    known_blocked = {
        f"i1: PDF embeds Rubik ({sorted(embedded)})": has_rubik,
        f"i2: PDF embeds {expected_body_font}": has_body_font,
        f"i3: PDF embeds no Calibri/TimesNewRoman/ArialMT (off-brand={off_brand})": not off_brand,
    }

    failed = [name for name, ok in checks.items() if not ok]
    for name, ok in checks.items():
        print(f"[qa {doc_type}-{lang}] {'OK ' if ok else 'FAIL'} - {name}")
    for name, ok in known_blocked.items():
        print(f"[qa {doc_type}-{lang}] {'OK ' if ok else 'KNOWN-BLOCKED (relogin required)'} - {name}")
    if failed:
        raise AssertionError(f"QA failed for {doc_type}-{lang}: {failed}")
    print(f"[qa {doc_type}-{lang}] all required acceptance checks passed"
          + ("" if all(known_blocked.values()) else " (i-checks pending relogin)"))
    return metrics, (not failed), known_blocked


if __name__ == "__main__":
    # Console-encoding hardening: check names embed Hebrew (needles, measured
    # text), and on a cp1252-piped stdout print() would raise
    # UnicodeEncodeError and kill an otherwise-green build (best case:
    # mojibake, as seen in the first full run). Display-only - AUDITED: every
    # acceptance comparison (_contains, row_extremes, tagline_text_ok, the
    # XML greps) operates on in-memory unicode str, never on console-encoded
    # bytes, so the console codec cannot affect pass/fail results.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    rasterize_svgs()

    results = []
    for doc_type in DOC_TYPES:
        for lang in LANGS:
            docx_path, pdf_path = out_paths(doc_type, lang)
            label = f"{doc_type}-{lang}"
            try:
                build_doc(doc_type, lang, out_path=docx_path)
                verify_docx(docx_path)
                export_pdf(docx_path, pdf_path)
                metrics, ok, known_blocked = qa_checks(doc_type, lang, docx_path, pdf_path)
                i_pending = not all(known_blocked.values())
                results.append((label, "PASS" if ok else "FAIL", i_pending, None))
            except Exception as e:
                results.append((label, "FAIL", None, str(e)))
                print(f"[BUILD FAILED] {label}: {e}")

    print("\n" + "=" * 72)
    print(f"{'file':<16}{'result':<8}{'i-checks':<24}{'error'}")
    print("=" * 72)
    for label, result, i_pending, err in results:
        i_status = "pending relogin" if i_pending else ("n/a" if i_pending is None else "OK")
        print(f"{label:<16}{result:<8}{i_status:<24}{err or ''}")
    print("=" * 72)

    n_fail = sum(1 for _, r, _, _ in results if r != "PASS")
    if n_fail:
        print(f"\n{n_fail} of {len(results)} builds FAILED.")
        sys.exit(1)
    print(f"\nAll {len(results)} builds complete and passed QA (i-checks pending relogin, see above).")
